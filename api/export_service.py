import io
import csv
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

class ExportService:
    @staticmethod
    def generate_csv(entries: List[Dict[str, Any]]) -> str:
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["Day", "Period", "Section", "Subject Code", "Teacher", "Room", "Type", "Group", "Locked"])
        for e in sorted(entries, key=lambda x: (x["day"], x["period"], x["section_id"])):
            writer.writerow([
                e["day"],
                e["period"],
                e["section_id"],
                e["subject_id"],
                e["teacher_id"],
                e["room_id"],
                "Elective" if e.get("is_elective") else "Regular",
                e.get("elective_group") or "-",
                "Yes" if e.get("is_locked") else "No"
            ])
        return output.getvalue()

    @staticmethod
    def generate_excel(
        entries: List[Dict[str, Any]],
        days: List[str],
        periods: List[int],
        teachers_map: Dict[str, Any],
        subjects_map: Dict[str, Any],
        sections: List[Dict[str, Any]],
        time_slots: Optional[List[Dict[str, Any]]] = None
    ) -> io.BytesIO:
        wb = openpyxl.Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        # Style tokens
        header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        
        day_header_font = Font(name="Segoe UI", size=10, bold=True, color="1E293B")
        day_header_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

        cell_font = Font(name="Segoe UI", size=10)
        cell_sub_font = Font(name="Segoe UI", size=9, bold=True, color="0F172A")
        cell_detail_font = Font(name="Segoe UI", size=8, color="475569")
        
        lab_fill = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid") # Light blue
        elec_fill = PatternFill(start_color="F3E8FF", end_color="F3E8FF", fill_type="solid") # Light purple
        regular_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        locked_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid") # Amber

        thin_side = Side(style="thin", color="CBD5E1")
        border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Build slot label map
        slot_label_map = {}
        if time_slots:
            for s in time_slots:
                if s.get("period", 0) > 0:
                    slot_label_map[s["period"]] = f"Period {s['period']}\n({s.get('start_time', '')} - {s.get('end_time', '')})"

        # 1. Sheets for each Section
        for sec in sections:
            sec_id = sec["id"]
            ws = wb.create_sheet(title=f"Sec {sec_id}"[:31])
            
            # Title
            ws.merge_cells("A1:H1")
            title_cell = ws["A1"]
            title_cell.value = f"Class Timetable: {sec_id} ({sec.get('branch', '')} - Semester {sec.get('semester', 5)})"
            title_cell.font = Font(name="Segoe UI", size=14, bold=True, color="0F172A")
            title_cell.alignment = Alignment(horizontal="left", vertical="center")
            ws.row_dimensions[1].height = 30

            # Headers (Row 3)
            headers = ["Day / Period"] + [slot_label_map.get(p, f"Period {p}") for p in periods]
            ws.row_dimensions[3].height = 32
            for col_idx, h in enumerate(headers, start=1):
                c = ws.cell(row=3, column=col_idx, value=h)
                c.font = header_font
                c.fill = header_fill
                c.alignment = center_align
                c.border = border

            # Fill Day rows
            row_idx = 4
            for day in days:
                ws.row_dimensions[row_idx].height = 42
                # Day label
                day_cell = ws.cell(row=row_idx, column=1, value=day)
                day_cell.font = day_header_font
                day_cell.fill = day_header_fill
                day_cell.alignment = center_align
                day_cell.border = border

                for p_idx, p in enumerate(periods, start=2):
                    slot_entries = [
                        e for e in entries
                        if e["section_id"] == sec_id and e["day"] == day and e["period"] == p
                    ]
                    c = ws.cell(row=row_idx, column=p_idx)
                    c.border = border
                    c.alignment = center_align

                    if not slot_entries:
                        c.value = "—"
                        c.fill = regular_fill
                    elif len(slot_entries) == 1:
                        e = slot_entries[0]
                        subj_name = subjects_map.get(e["subject_id"], {}).get("name", e["subject_id"])
                        t_name = teachers_map.get(e["teacher_id"], {}).get("name", e["teacher_id"])
                        is_lab = subjects_map.get(e["subject_id"], {}).get("needs_lab", False)
                        
                        c.value = f"{e['subject_id']}\n{t_name} | {e['room_id']}"
                        if e.get("is_locked"):
                            c.fill = locked_fill
                        elif is_lab:
                            c.fill = lab_fill
                        elif e.get("is_elective"):
                            c.fill = elec_fill
                        else:
                            c.fill = regular_fill
                    else:
                        # Parallel Electives
                        lines = []
                        for e in slot_entries:
                            lines.append(f"[{e['subject_id']}] {e['teacher_id']}/{e['room_id']}")
                        c.value = "⚡ " + (slot_entries[0].get("elective_group") or "Elective") + "\n" + "\n".join(lines)
                        c.fill = elec_fill

                row_idx += 1

            # Adjust column widths
            ws.column_dimensions["A"].width = 14
            for p_idx in range(2, len(periods) + 2):
                col_letter = get_column_letter(p_idx)
                ws.column_dimensions[col_letter].width = 20

        # 2. Master Sheet (All Sections)
        ws_all = wb.create_sheet(title="Master Schedule")
        ws_all.merge_cells("A1:F1")
        m_title = ws_all["A1"]
        m_title.value = "Master College Timetable Summary"
        m_title.font = Font(name="Segoe UI", size=14, bold=True, color="0F172A")
        ws_all.row_dimensions[1].height = 28

        m_headers = ["Day", "Period", "Section", "Subject", "Teacher", "Room", "Type", "Status"]
        ws_all.row_dimensions[3].height = 24
        for col_idx, h in enumerate(m_headers, start=1):
            c = ws_all.cell(row=3, column=col_idx, value=h)
            c.font = header_font
            c.fill = header_fill
            c.alignment = center_align
            c.border = border

        m_row = 4
        for e in sorted(entries, key=lambda x: (x["day"], x["period"], x["section_id"])):
            subj_name = subjects_map.get(e["subject_id"], {}).get("name", e["subject_id"])
            t_name = teachers_map.get(e["teacher_id"], {}).get("name", e["teacher_id"])
            type_str = f"Elective ({e.get('elective_group', '')})" if e.get("is_elective") else ("Lab" if subjects_map.get(e["subject_id"], {}).get("needs_lab") else "Theory")
            
            row_vals = [
                e["day"],
                f"P{e['period']}",
                e["section_id"],
                f"{e['subject_id']} ({subj_name})",
                f"{e['teacher_id']} - {t_name}",
                e["room_id"],
                type_str,
                "🔒 Locked" if e.get("is_locked") else "Auto-assigned"
            ]
            for col_idx, val in enumerate(row_vals, start=1):
                c = ws_all.cell(row=m_row, column=col_idx, value=val)
                c.font = cell_font
                c.border = border
                if col_idx in [1, 2, 3, 6, 7, 8]:
                    c.alignment = center_align
            m_row += 1

        for col_idx, width in enumerate([10, 10, 14, 34, 28, 12, 18, 16], start=1):
            ws_all.column_dimensions[get_column_letter(col_idx)].width = width

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output
